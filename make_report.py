#!/usr/bin/env python3
"""Build the hourly machine time workbook from Daily Production Report PDFs.

Usage:
  python make_report.py PDF_FOLDER [-o out.xlsx] [--threshold 10] [--expect "B650 T" "B1000 T"]
Needs: pip install pdfplumber openpyxl
"""
import argparse, glob, os, re, sys
from datetime import timedelta
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule

ROW_RE = re.compile(r"^\d{1,2}:\d{2}\s*-\s*\d{1,2}:\d{2}$")

def dur(s):
    s = (s or "").strip()
    g = lambda u: int(m.group(1)) if (m := re.search(r"(\d+)\s*" + u, s)) else 0
    return g("h") * 3600 + g("m") * 60 + g("s")

def norm_name(n):
    return re.sub(r"\s+", " ", re.sub(r"\(", " (", n)).strip()

def parse_pdf(path):
    machine = date = None
    rows = []
    with pdfplumber.open(path) as pdf:
        for pg in pdf.pages:
            txt = pg.extract_text() or ""
            if machine is None:
                m = re.search(r"Daily Production Report\s+(.+?)\s*\|", txt)
                if m: machine = norm_name(m.group(1))
                d = re.search(r"(\d{1,2} [A-Za-z]{3} \d{4})", txt)
                if d: date = d.group(1)
            for tb in pg.extract_tables():
                for r in tb:
                    if r and r[0] and ROW_RE.match(r[0].strip()) and len(r) >= 9:
                        rows.append(dict(time=r[0].strip(), status=r[1], shots=r[2],
                                         on=dur(r[6]), idle=dur(r[7]), off=dur(r[8])))
    return machine, date, rows

def build_workbook(data, date, out, threshold=10, expect=()):
    order = sorted(set(data) | set(expect))
    slots = [(9 + i) % 24 for i in range(24)]

    F = "Arial"
    hf = Font(name=F, bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="1F3864")
    body = Font(name=F); blue = Font(name=F, color="0000FF")
    thin = Side(style="thin", color="BFBFBF"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    DUR = "[h]:mm:ss"

    wb = Workbook(); ws = wb.active; ws.title = "Hourly"
    ws["A1"] = f"Hourly Machine Time Report - {date} 09:00 to next day 09:00"; ws["A1"].font = Font(name=F, bold=True, size=13)
    ws["A2"] = "Threshold:"; ws["A2"].font = Font(name=F, bold=True)
    ws["B2"] = f"=TIME(0,{threshold},0)"; ws["B2"].number_format = DUR; ws["B2"].font = Font(name=F, bold=True, color="0000FF")
    heads = ["Time Duration", "Shift", "Machine Name", "On Time", "Idle Time", "Off Time", "Missing Time", f"Above {threshold} Min Threshold"]
    for c, h in enumerate(heads, 1):
        x = ws.cell(3, c, h); x.font = hf; x.fill = hfill; x.border = bd; x.alignment = Alignment(horizontal="center")
    r = 4
    for m in order:
        if m not in data: continue
        agg = {s: [0, 0, 0] for s in slots}
        for row in data[m]:
            h = int(row["time"].split(":")[0])
            agg[h][0] += row["on"]; agg[h][1] += row["idle"]; agg[h][2] += row["off"]
        for s in slots:
            if sum(agg[s]) == 0: agg[s][2] = 3600   # no activity at all in this hour = machine off
            ws.cell(r, 1, f"{s:02d}:00 - {(s + 1) % 24:02d}:00")
            ws.cell(r, 2, "A (09:00-21:00)" if 9 <= s < 21 else "B (21:00-09:00)")
            ws.cell(r, 3, m)
            for c, v in zip((4, 5, 6), agg[s]):
                ws.cell(r, c, timedelta(seconds=v)).font = blue
            ws.cell(r, 7, f"=MAX(0,TIME(1,0,0)-(D{r}+E{r}+F{r}))")
            ws.cell(r, 8, f'=IF(G{r}>$B$2,"Above threshold","Within threshold")').alignment = Alignment(horizontal="center")
            for c in range(1, 9):
                ws.cell(r, c).border = bd
                if c not in (4, 5, 6): ws.cell(r, c).font = body
            for c in (4, 5, 6, 7): ws.cell(r, c).number_format = DUR
            r += 1
    last = r - 1
    for c, w in zip("ABCDEFGH", [16, 16, 16, 12, 12, 12, 14, 22]): ws.column_dimensions[c].width = w
    ws.freeze_panes = "A4"; ws.auto_filter.ref = f"A3:H{last}"
    red = PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE"); grn = PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")
    ws.conditional_formatting.add(f"G4:H{last}", FormulaRule(formula=["$G4>$B$2"], fill=red, font=Font(name=F, color="9C0006")))
    ws.conditional_formatting.add(f"G4:H{last}", FormulaRule(formula=["$G4<=$B$2"], fill=grn, font=Font(name=F, color="006100")))
    n = last + 2
    ws.cell(n, 1, "Notes").font = Font(name=F, bold=True)
    for i, t in enumerate(["On, Idle and Off Time (blue) = sum of report rows that start in that clock hour.",
                           "An hour with no activity at all (no report rows) is counted as Off Time = 1:00:00, so Missing Time is 0.",
                           "Missing Time = 1 hour - (On + Idle + Off), never below 0."]):
        ws.cell(n + 1 + i, 1, t).font = body

    sm = wb.create_sheet("Summary")
    for c, h in enumerate(["Machine Name", "On Time", "Idle Time", "Off Time", "Missing Time", "Note"], 1):
        x = sm.cell(1, c, h); x.font = hf; x.fill = hfill; x.border = bd; x.alignment = Alignment(horizontal="center")
    for i, m in enumerate(order, 2):
        sm.cell(i, 1, m)
        if m in data:
            for c, col in ((2, "D"), (3, "E"), (4, "F"), (5, "G")):
                sm.cell(i, c, f"=SUMIFS(Hourly!${col}$4:${col}${last},Hourly!$C$4:$C${last},$A{i})").number_format = DUR
        else:
            sm.cell(i, 6, "No data in report")
        for c in range(1, 7): sm.cell(i, c).font = body; sm.cell(i, c).border = bd
    for c, w in zip("ABCDEF", [16, 14, 14, 14, 14, 22]): sm.column_dimensions[c].width = w
    wb.save(out)



def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("folder"); ap.add_argument("-o", "--out")
    ap.add_argument("--threshold", type=int, default=10, help="minutes")
    ap.add_argument("--expect", nargs="*", default=[], help="machines that should appear")
    a = ap.parse_args()

    data, date = {}, None
    for p in sorted(glob.glob(os.path.join(a.folder, "*.pdf"))):
        mach, d, rows = parse_pdf(p)
        if not mach: print("skip (no machine name):", p); continue
        date = date or d
        seen = data.setdefault(mach, [])
        for r in rows:
            if r not in seen: seen.append(r)   # same row in two shift PDFs counts once
        print(f"{mach}: {len(rows)} rows from {os.path.basename(p)}")
    if not data: sys.exit("No reports found")
    out = a.out or f"Hourly_Machine_Time_Report_{(date or 'report').replace(' ', '')}.xlsx"
    build_workbook(data, date, out, a.threshold, a.expect); print("Saved", out)


if __name__ == "__main__":
    main()


def hourly_table(data, order=None):
    """Return a list of dict rows (seconds) for every machine and clock hour."""
    slots = [(9 + i) % 24 for i in range(24)]
    out = []
    for m in (order or sorted(data)):
        if m not in data: continue
        agg = {s: [0, 0, 0] for s in slots}
        for row in data[m]:
            h = int(row["time"].split(":")[0])
            agg[h][0] += row["on"]; agg[h][1] += row["idle"]; agg[h][2] += row["off"]
        for s in slots:
            on, idle, off = agg[s]
            if on + idle + off == 0: off = 3600   # no activity at all in this hour = machine off
            out.append(dict(hour=f"{s:02d}:00 - {(s + 1) % 24:02d}:00",
                            shift="A" if 9 <= s < 21 else "B", machine=m,
                            on=on, idle=idle, off=off, missing=max(0, 3600 - (on + idle + off))))
    return out


# ---------- Pasted dashboard text ----------
_SPAN_RE = re.compile(r"(\d{1,2}):(\d{2})\s*-\s*(\d{1,2}):(\d{2})")
_DUR_RE = re.compile(r"(?:\d+h[ ]*)?(?:\d+m[ ]*)?(?:\d+s)?")


def parse_dashboard_text(text):
    """Parse hourly rows copied from the dashboard's 'Hourly Details' table.

    Works with tab/line-separated paste and with text that was pasted with no
    separators at all. Returns rows in the same format as parse_pdf().
    Whole-hour rows (09:00 - 10:00) are used when present; otherwise the
    smaller rows inside the hour are used. Shift rows (09:00 - 21:00) are ignored.
    """
    t = text.replace("\r", "").replace("\t", "|").replace("\n", "|")
    spans = list(_SPAN_RE.finditer(t))
    hours, subs = {}, []
    for i, m in enumerate(spans):
        seg = t[m.end(): spans[i + 1].start() if i + 1 < len(spans) else len(t)]
        sm = re.search(r"(\d+)\s*shot\(s\)", seg)
        km = re.search(r"kWh", seg)
        if not sm or not km:
            continue
        pos, toks = km.end(), []
        while len(toks) < 4:
            while pos < len(seg) and seg[pos] in " |":
                pos += 1
            d = _DUR_RE.match(seg, pos)
            if not d or d.end() == pos:
                break
            toks.append(dur(d.group())); pos = d.end()
        if len(toks) < 3:
            continue
        h1, m1, h2, m2 = (int(x) for x in m.groups())
        start, end = h1 * 60 + m1, h2 * 60 + m2
        if end <= start:
            end += 1440
        span = end - start
        if span >= 240:          # shift / total rows
            continue
        row = dict(time=f"{h1:02d}:{m1:02d} - {h2:02d}:{m2:02d}",
                   status=seg[:sm.start()].strip(" |,"), shots=f"{sm.group(1)} shot(s)",
                   on=toks[0], idle=toks[1], off=toks[2])
        if start % 60 == 0 and span == 60:
            hours.setdefault(h1 % 24, row)
        elif row not in subs:
            subs.append(row)
    return list(hours.values()) + [r for r in subs if int(r["time"][:2]) % 24 not in hours]
